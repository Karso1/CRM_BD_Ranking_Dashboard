import importlib.util
from datetime import date, datetime, timedelta
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from unittest.mock import MagicMock

SPEC = importlib.util.spec_from_file_location('upb_download', Path(__file__).resolve().parents[1] / 'tools/upb-download/download.py')
MODULE = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(MODULE)


class BatchTests(unittest.TestCase):
    def test_manual_page_count_reads_only_pagination(self):
        page = MagicMock()
        total = page.locator.return_value
        total.count.return_value = 1
        total.inner_text.return_value = '共 3633 条'
        self.assertEqual(MODULE.manual_page_count(page), 3633)
        page.locator.assert_called_once_with('main .pagination-container > p')
        total.inner_text.return_value = '共 4864 条；其他统计'
        with self.assertRaisesRegex(RuntimeError, '格式已变化'):
            MODULE.manual_page_count(page)

    def test_manual_search_allows_unchanged_result_count(self):
        page = MagicMock()
        page.locator.return_value.locator.return_value.inner_text.return_value = MODULE.MANUAL_STATUS
        with patch.object(MODULE, 'wait_idle'), patch.object(MODULE, 'wait_manual_controls_ready'), \
             patch.object(MODULE, 'manual_result_count_after_search', return_value=3633):
            self.assertEqual(MODULE.prepare_manual_page(page), 3633)

    def test_manual_search_retries_transient_zero_without_count_change_requirement(self):
        page = MagicMock()
        with patch.object(MODULE, 'wait_manual_controls_ready'), \
             patch.object(MODULE, 'manual_page_count', side_effect=[0, 3633]), \
             patch.object(MODULE, 'manual_visible_rows_success', return_value=True):
            self.assertEqual(MODULE.manual_result_count_after_search(page), 3633)
        self.assertEqual(page.get_by_role.return_value.click.call_count, 2)

    def test_manual_search_three_zero_results_stops_safely(self):
        page = MagicMock()
        with patch.object(MODULE, 'wait_manual_controls_ready'), \
             patch.object(MODULE, 'manual_page_count', return_value=0):
            with self.assertRaisesRegex(MODULE.PageNotReady, '连续 3 次搜索'):
                MODULE.manual_result_count_after_search(page)
        self.assertEqual(page.get_by_role.return_value.click.call_count, 3)

    def test_manual_search_checks_all_visible_status_cells(self):
        page = MagicMock()
        header = MagicMock()
        header.inner_text.return_value = '卡片状态'
        good = MagicMock()
        good.locator.return_value.nth.return_value.inner_text.return_value = '绑定成功'
        pending = MagicMock()
        pending.locator.return_value.nth.return_value.inner_text.return_value = '待送出'
        page.locator.side_effect = lambda selector: MagicMock(all=lambda: [header] if 'thead' in selector else [good, pending])
        self.assertFalse(MODULE.manual_visible_rows_success(page))

    def manual_workbook(self, path, status='绑定成功', card_id='card-1'):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.append(['历史代理商', '用户卡片ID', '渠道卡ID', '申请日期', '订单状态'])
        sheet.append(['UPay', card_id, 'channel-1', '2026-10-18', status])
        book.save(path)

    def test_manual_snapshot_accepts_success_and_rejects_other_status(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'manual.xlsx'
            self.manual_workbook(path)
            self.assertEqual(MODULE.validate_manual(path), 1)
            self.manual_workbook(path, status='处理中')
            with self.assertRaisesRegex(ValueError, '不是成功状态'):
                MODULE.validate_manual(path)
            self.manual_workbook(path, card_id='')
            with self.assertRaisesRegex(ValueError, '用户卡片ID'):
                MODULE.validate_manual(path)

    def test_manual_snapshot_replaces_only_file_and_keeps_backup(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            data = root / 'data'
            folder = data / '手动开卡'
            folder.mkdir(parents=True)
            state = root / 'state'
            state.mkdir()
            old = folder / '手动开卡旧日期.xlsx'
            self.manual_workbook(old, card_id='old-card')
            original_hash = MODULE.digest(old)
            fresh = root / 'new.xlsx'
            self.manual_workbook(fresh, card_id='new-card')
            target, backup = MODULE.replace_manual_snapshot(fresh, data, state, 1, 'export.xlsx')
            self.assertEqual(target, folder / '手动开卡-最新.xlsx')
            self.assertEqual(MODULE.digest(target), MODULE.digest(fresh))
            self.assertEqual(MODULE.digest(backup), original_hash)
            self.assertFalse(old.exists())
            self.assertEqual(len(list(folder.glob('*.xlsx'))), 1)
            record = json.loads((state / 'manual-snapshot.json').read_text())
            self.assertEqual(record['rows'], 1)
            self.assertEqual(record['path'], str(target))
            next_export = root / 'next.xlsx'
            self.manual_workbook(next_export, card_id='next-card')
            next_target, next_backup = MODULE.replace_manual_snapshot(next_export, data, state, 1, 'next.xlsx')
            self.assertEqual(next_target, target)
            self.assertEqual(MODULE.digest(next_backup), MODULE.digest(fresh))
            self.assertEqual(len(list(folder.glob('*.xlsx'))), 1)

    def test_manual_snapshot_multiple_originals_are_not_replaced(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            folder = root / 'data' / '手动开卡'
            folder.mkdir(parents=True)
            state = root / 'state'
            state.mkdir()
            self.manual_workbook(folder / 'one.xlsx')
            self.manual_workbook(folder / 'two.xlsx')
            fresh = root / 'fresh.xlsx'
            self.manual_workbook(fresh)
            before = {p.name: MODULE.digest(p) for p in folder.glob('*.xlsx')}
            with self.assertRaisesRegex(RuntimeError, '多份 Excel'):
                MODULE.replace_manual_snapshot(fresh, root / 'data', state, 1, 'fresh.xlsx')
            self.assertEqual({p.name: MODULE.digest(p) for p in folder.glob('*.xlsx')}, before)

    def test_manual_snapshot_migration_failure_restores_legacy_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            folder = root / 'data' / '手动开卡'
            folder.mkdir(parents=True)
            state = root / 'state'
            state.mkdir()
            old = folder / '手动开卡旧日期.xlsx'
            self.manual_workbook(old, card_id='old-card')
            original_hash = MODULE.digest(old)
            fresh = root / 'fresh.xlsx'
            self.manual_workbook(fresh, card_id='new-card')
            replace = MODULE.os.replace
            def fail_manifest(source, target):
                if str(target).endswith('manual-snapshot.json'):
                    raise OSError('simulated manifest failure')
                return replace(source, target)
            with patch.object(MODULE.os, 'replace', side_effect=fail_manifest):
                with self.assertRaisesRegex(OSError, 'simulated manifest failure'):
                    MODULE.replace_manual_snapshot(fresh, root / 'data', state, 1, 'fresh.xlsx')
            self.assertEqual(MODULE.digest(old), original_hash)
            self.assertFalse((folder / '手动开卡-最新.xlsx').exists())

    def setup_batch(self, root):
        batch = root / 'batch'
        data = root / 'data'
        batch.mkdir()
        data.mkdir()
        for key in MODULE.SOURCES:
            (batch / f'{key}.xlsx').write_bytes(key.encode())
        return batch, data, {key: {} for key in MODULE.SOURCES}

    def test_existing_file_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            batch, data, records = self.setup_batch(Path(tmp))
            dest = data / '消费数据/每日自动下载/2026-10/Reap_2026-10-15.xlsx'
            dest.parent.mkdir(parents=True)
            dest.write_bytes(b'original')
            with self.assertRaisesRegex(RuntimeError, '已存在'):
                MODULE.commit_batch(batch, date(2026, 10, 15), records, data)
            self.assertEqual(dest.read_bytes(), b'original')
            self.assertEqual(len(list(data.rglob('*.xlsx'))), 1)

    def test_copy_failure_rolls_back_and_keeps_downloads(self):
        with tempfile.TemporaryDirectory() as tmp:
            batch, data, records = self.setup_batch(Path(tmp))
            copy = MODULE.shutil.copy2
            def fail_third(source, target):
                if source.name == 'passto.xlsx':
                    raise OSError('simulated disk failure')
                return copy(source, target)
            with patch.object(MODULE.shutil, 'copy2', side_effect=fail_third):
                with self.assertRaises(OSError):
                    MODULE.commit_batch(batch, date(2026, 10, 15), records, data)
            self.assertEqual(list(data.rglob('*.xlsx')), [])
            self.assertEqual(len(list(batch.glob('*.xlsx'))), 5)

    def test_success_records_exact_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            batch, data, records = self.setup_batch(Path(tmp))
            MODULE.commit_batch(batch, date(2026, 10, 15), records, data)
            self.assertEqual(len(list(data.rglob('*.xlsx'))), 5)
            for key, record in records.items():
                self.assertEqual(Path(record['path']).read_bytes(), key.encode())
                self.assertEqual(record['sha256'], MODULE.digest(Path(record['path'])))

    def test_repair_keeps_backup_and_only_replaces_channels(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            batch, data, records = self.setup_batch(root)
            day = date(2026, 10, 15)
            MODULE.commit_batch(batch, day, records, data)
            (root / f'{day}.json').write_text(json.dumps({'date': str(day), 'files': records}))
            old = {source: Path(record['path']).read_bytes() for source, record in records.items()}
            for source in MODULE.CHANNEL_SOURCES:
                (batch / f'{source}.xlsx').write_bytes(f'new-{source}'.encode())
                records[source]['rows'] = 1
            with patch.object(MODULE, 'validate', return_value=1):
                backup = MODULE.repair_batch(batch, day, records, data, root)
            for source, record in records.items():
                target = Path(record['path'])
                if source in MODULE.CHANNEL_SOURCES:
                    self.assertEqual(target.read_bytes(), f'new-{source}'.encode())
                    self.assertEqual((backup / target.name).read_bytes(), old[source])
                else:
                    self.assertEqual(target.read_bytes(), old[source])
            self.assertEqual(json.loads((root / f'{day}.json').read_text())['repair_backup'], str(backup))

    def test_repair_failure_restores_original_files_and_manifest(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            batch, data, records = self.setup_batch(root)
            day = date(2026, 10, 15)
            MODULE.commit_batch(batch, day, records, data)
            manifest = root / f'{day}.json'
            manifest.write_text(json.dumps({'date': str(day), 'files': records}))
            original_manifest = manifest.read_bytes()
            originals = {source: Path(record['path']).read_bytes() for source, record in records.items()}
            for source in MODULE.CHANNEL_SOURCES:
                (batch / f'{source}.xlsx').write_bytes(f'new-{source}'.encode())
                records[source]['rows'] = 1
            replace = MODULE.os.replace
            staged_replacements = 0
            def fail_second_replacement(source, target):
                nonlocal staged_replacements
                if str(source).endswith('.repair-pending'):
                    staged_replacements += 1
                    if staged_replacements == 2:
                        raise OSError('simulated replacement failure')
                return replace(source, target)
            with patch.object(MODULE, 'validate', return_value=1), \
                 patch.object(MODULE.os, 'replace', side_effect=fail_second_replacement):
                with self.assertRaisesRegex(OSError, 'simulated replacement failure'):
                    MODULE.repair_batch(batch, day, records, data, root)
            self.assertEqual(manifest.read_bytes(), original_manifest)
            for source, old_bytes in originals.items():
                self.assertEqual(Path(records[source]['path']).read_bytes(), old_bytes)

    def test_catch_up_includes_missed_days_across_month_boundary(self):
        with tempfile.TemporaryDirectory() as tmp:
            state = Path(tmp)
            (state / '2026-09-29.json').write_text('{}')
            days = MODULE.dates_to_download(None, None, None, date(2026, 10, 3), state)
            self.assertEqual(days, [date(2026, 9, 29), date(2026, 9, 30),
                                    date(2026, 10, 1), date(2026, 10, 2)])

    def test_first_run_uses_yesterday(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(MODULE.dates_to_download(None, None, None, date(2026, 10, 16), Path(tmp)),
                             [date(2026, 10, 15)])

    def test_explicit_range_and_invalid_dates(self):
        with tempfile.TemporaryDirectory() as tmp:
            state = Path(tmp)
            self.assertEqual(MODULE.dates_to_download(None, date(2026, 10, 12), date(2026, 10, 14),
                                                      date(2026, 10, 16), state),
                             [date(2026, 10, 12), date(2026, 10, 13), date(2026, 10, 14)])
            with self.assertRaises(ValueError):
                MODULE.dates_to_download(None, date(2026, 10, 14), None, date(2026, 10, 16), state)
            with self.assertRaises(ValueError):
                MODULE.dates_to_download(date(2026, 10, 16), None, None, date(2026, 10, 16), state)

    def test_completed_batch_requires_all_five_unchanged_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            state = Path(tmp)
            batch, data, records = self.setup_batch(state)
            MODULE.commit_batch(batch, date(2026, 10, 15), records, data)
            manifest = state / '2026-10-15.json'
            manifest.write_text(json.dumps({'date': '2026-10-15', 'files': records}))
            self.assertTrue(MODULE.completed_batch(date(2026, 10, 15), state))
            (Path(records['cards']['path'])).write_bytes(b'changed')
            with self.assertRaisesRegex(RuntimeError, '发生变化'):
                MODULE.completed_batch(date(2026, 10, 15), state)

    def test_split_window_has_no_overlap_or_gap(self):
        start = datetime(2026, 10, 15)
        end = start + timedelta(days=1) - timedelta(seconds=1)
        left, right = MODULE.split_window(start, end)
        self.assertEqual(left[0], start)
        self.assertEqual(left[1] + timedelta(seconds=1), right[0])
        self.assertEqual(right[1], end)

    def test_large_day_is_split_below_safe_export_count(self):
        start = datetime(2026, 10, 15)
        end = start + timedelta(days=1) - timedelta(seconds=1)
        intervals = []
        def prepare(_page, _source, first, last):
            intervals.append((first, last))
            return int((last - first).total_seconds()) + 1
        def export(_page, _source, path):
            path.touch()
            return path.name, int((intervals[-1][1] - intervals[-1][0]).total_seconds()) + 1
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(MODULE, 'prepare_page', side_effect=prepare), \
                 patch.object(MODULE, 'export_file', side_effect=export), \
                 patch.object(MODULE, 'validate', side_effect=lambda *args: intervals[-1][1].timestamp() - intervals[-1][0].timestamp() + 1):
                parts, total = MODULE.download_window(None, 'passto', start.date(), start, end, Path(tmp))
        self.assertEqual(total, 86400)
        self.assertEqual(sum(count for _, count in parts), total)
        self.assertTrue(all(count < MODULE.SAFE_EXPORT_ROWS for _, count in parts))
        self.assertGreater(len(parts), 1)

    def test_merge_parts_preserves_all_rows(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            parts = []
            for index, hour in enumerate((1, 2)):
                path = root / f'part{index}.xlsx'
                book = Workbook()
                sheet = book.active
                sheet.append(['更新时间', '订单类型', '交易金额'])
                sheet.append([datetime(2026, 10, 15, hour), '主订单', 10 + index])
                book.save(path)
                parts.append((path, 1))
            destination = root / 'merged.xlsx'
            self.assertEqual(MODULE.merge_parts('passto', date(2026, 10, 15), parts, destination), 2)
            self.assertEqual(MODULE.validate(destination, 'passto', date(2026, 10, 15)), 2)

    def test_channel_uses_update_date_and_rejects_child_orders(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'passto.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.append(['交易时间', '更新时间', '订单类型', '交易金额'])
            sheet.append([datetime(2026, 10, 14, 23), datetime(2026, 10, 15, 1), '主订单', 5])
            book.save(path)
            self.assertEqual(MODULE.validate(path, 'passto', date(2026, 10, 15)), 1)
            sheet['C2'] = '子订单'
            book.save(path)
            with self.assertRaisesRegex(ValueError, '不是主订单'):
                MODULE.validate(path, 'passto', date(2026, 10, 15))

    def test_straitsx_uses_transaction_day_and_requires_main_order(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'straitsx.xlsx'
            book = Workbook()
            sheet = book.active
            sheet.append(['订单类型', '交易时间'])
            sheet.append(['主订单', datetime(2026, 10, 15, 23)])
            book.save(path)
            self.assertEqual(MODULE.validate(path, 'straitsx', date(2026, 10, 15)), 1)
            sheet['A2'] = '子订单'
            book.save(path)
            with self.assertRaisesRegex(ValueError, '不是主订单'):
                MODULE.validate(path, 'straitsx', date(2026, 10, 15))

    def test_stale_page_total_does_not_cause_repeated_splits(self):
        start = datetime(2026, 10, 15)
        end = start + timedelta(seconds=1)
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(MODULE, 'prepare_page', return_value=99999), \
                 patch.object(MODULE, 'export_file', return_value=('part.xlsx', 2)) as exported, \
                 patch.object(MODULE, 'validate', return_value=2):
                parts, total = MODULE.download_window(None, 'reap', start.date(), start, end, Path(tmp))
            self.assertEqual(total, 2)
            self.assertEqual(len(parts), 1)
            exported.assert_called_once()

    def test_task_and_file_row_mismatch_is_rejected(self):
        start = datetime(2026, 10, 15)
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(MODULE, 'prepare_page', return_value=2), \
                 patch.object(MODULE, 'export_file', return_value=('part.xlsx', 3)), \
                 patch.object(MODULE, 'validate', return_value=2):
                with self.assertRaisesRegex(RuntimeError, '文件实际 2 行'):
                    MODULE.download_window(None, 'reap', start.date(), start, start, Path(tmp))

    def test_missing_export_reopens_same_date_once(self):
        start = datetime(2026, 10, 21)
        end = start + timedelta(days=1) - timedelta(seconds=1)
        page = MagicMock()
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(MODULE, 'prepare_page', return_value=1775) as prepared, \
                 patch.object(MODULE, 'export_file', side_effect=[MODULE.ExportButtonMissing('missing'), ('part.xlsx', 1775)]) as exported, \
                 patch.object(MODULE, 'validate', return_value=1775):
                _, total = MODULE.download_window(page, 'recharge', start.date(), start, end, Path(tmp))
        self.assertEqual(total, 1775)
        self.assertEqual(prepared.call_count, 2)
        self.assertEqual(exported.call_count, 2)
        self.assertEqual(prepared.call_args_list[0].args, prepared.call_args_list[1].args)
        page.reload.assert_called_once_with(wait_until='domcontentloaded')

    def test_missing_export_twice_stops_without_installing(self):
        start = datetime(2026, 10, 21)
        page = MagicMock()
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(MODULE, 'prepare_page', return_value=1775) as prepared, \
                 patch.object(MODULE, 'export_file', side_effect=MODULE.ExportButtonMissing('missing')) as exported:
                with self.assertRaises(MODULE.ExportButtonMissing):
                    MODULE.download_window(page, 'recharge', start.date(), start, start, Path(tmp))
        self.assertEqual(prepared.call_count, 2)
        self.assertEqual(exported.call_count, 2)
        page.reload.assert_called_once_with(wait_until='domcontentloaded')

    def test_stuck_loading_button_is_not_treated_as_a_finished_search(self):
        page = MagicMock()
        page.locator.return_value.wait_for.side_effect = TimeoutError('still spinning')
        with self.assertRaisesRegex(MODULE.PageNotReady, '仍在加载'):
            MODULE.wait_controls_ready(page, '充值 搜索')
        page.locator.assert_called_once_with('main .ant-btn-loading')

    def test_stuck_search_refreshes_and_reapplies_same_date(self):
        start = datetime(2026, 10, 24)
        page = MagicMock()
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(MODULE, 'prepare_page', side_effect=[MODULE.PageNotReady('stuck'), 1518]) as prepared, \
                 patch.object(MODULE, 'export_file', return_value=('part.xlsx', 1518)) as exported, \
                 patch.object(MODULE, 'validate', return_value=1518):
                _, total = MODULE.download_window(page, 'recharge', start.date(), start, start, Path(tmp))
        self.assertEqual(total, 1518)
        self.assertEqual(prepared.call_count, 2)
        self.assertEqual(exported.call_count, 1)
        self.assertEqual(prepared.call_args_list[0].args, prepared.call_args_list[1].args)
        page.reload.assert_called_once_with(wait_until='domcontentloaded')

    def test_export_button_ignores_history_action(self):
        page = MagicMock()
        history = MagicMock()
        history.inner_text.return_value = '导出历史'
        export = MagicMock()
        export.inner_text.return_value = 'Excel 导出'
        page.get_by_role.return_value.filter.return_value.all.return_value = [history, export]
        self.assertIs(MODULE.find_export_button(page, 'recharge'), export)
        page.evaluate.assert_called_once()

    def test_missing_export_reports_available_controls(self):
        page = MagicMock()
        history = MagicMock()
        history.inner_text.return_value = '导出历史'
        page.get_by_role.return_value.filter.return_value.all.return_value = [history]
        with patch.object(MODULE.time, 'monotonic', side_effect=[0, 13]):
            with self.assertRaisesRegex(MODULE.ExportButtonMissing, '当前可见按钮.*导出历史'):
                MODULE.find_export_button(page, 'recharge')

    def test_delayed_export_history_chooses_newest_task(self):
        before = {'2111688838963000672'}
        current = before | {'2111688891755810816', '2111688900505128960'}
        self.assertEqual(MODULE.select_new_task(before, current), '2111688900505128960')
        self.assertIsNone(MODULE.select_new_task(current, current))
