"""Create the compact UPay Wallet dashboard dataset.

Usage: python3 scripts/import_uw_excel.py /path/to/UW每日数据.xlsx app/wallet-data.json
"""
from __future__ import annotations

import calendar
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


MONTHS = {
    1: ("UW maintainer汇总1", "UW代理日汇总数据1"),
    2: ("UW maintainer汇总2", "UW代理日汇总数据2"),
    3: ("汇总33", "uw代理商日汇总数据3"),
    4: ("汇总4", "UW代理商日汇总4"),
    5: ("汇总5", "代理日汇总5"),
    6: ("汇总6", "代理日汇总6"),
    7: ("汇总7", "代理日汇总7"),
    8: ("汇总8", "代理日汇总8"),
    9: ("汇总9", "代理日汇总9"),
}


def number(value):
    return float(value or 0) if isinstance(value, (int, float)) else 0.0


def clean(value):
    return str(value).strip() if value not in (None, "") else ""


def valid_date(value):
    try:
        datetime.strptime(value, "%Y-%m-%d")
        return True
    except ValueError:
        return False


def parse_daily_reports(sheet):
    """Read the repeated daily blocks and retain the last duplicate date."""
    reports = {}
    current_date, current = None, {}
    title = re.compile(r"Performance Analysis\s+(\d{4}/\d{2}/\d{2})")

    def save():
        if current_date and valid_date(current_date):
            reports[current_date] = list(current.values())

    for row in sheet.iter_rows(values_only=True):
        match = title.search(clean(row[0] if row else None))
        if match:
            save()
            current_date = match.group(1).replace("/", "-")
            current = {}
            continue
        if not current_date:
            continue
        owner, name = clean(row[0]), clean(row[1])
        if not owner or not name or owner.lower() in {"maintainer", "total"}:
            continue
        key = (owner.upper(), name)
        current[key] = {
            "name": name,
            "owner": owner.upper(),
            "type": "代理商",
            "recharge": number(row[5] if len(row) > 5 else 0),
            "consumption": number(row[5] if len(row) > 5 else 0),
            "cards": number(row[3] if len(row) > 3 else 0) + number(row[4] if len(row) > 4 else 0),
        }
    save()
    return [{"date": date, "details": reports[date]} for date in sorted(reports)]


def final_details(sheet, last_daily):
    """Read cumulative values from the final daily block's right-hand table."""
    rows = list(sheet.iter_rows(values_only=True))
    starts = [i for i, row in enumerate(rows) if "Performance Analysis" in clean(row[0] if row else None)]
    block = rows[starts[-1]:] if starts else rows
    daily_lookup = {(row["owner"], row["name"]): row for row in last_daily}
    details = []
    for row in block[2:]:
        owner, name = clean(row[0]), clean(row[1])
        if not owner or not name or owner.lower() in {"maintainer", "total"}:
            continue
        key = (owner.upper(), name)
        cumulative = number(row[12] if len(row) > 12 else 0)
        cards = number(row[10] if len(row) > 10 else 0) + number(row[11] if len(row) > 11 else 0)
        details.append({
            "name": name,
            "owner": key[0],
            "type": "代理商",
            "recharge": cumulative,
            "consumption": cumulative,
            "cards": cards,
            "yesterday": daily_lookup.get(key, {}).get("recharge", 0),
        })
    return details


def parse_overall(sheet, details):
    rows = list(sheet.iter_rows(values_only=True))
    header = next((i for i, row in enumerate(rows) if clean(row[0]).lower() == "maintainer" and clean(row[1]).lower() == "target"), None)
    if header is None:
        return []
    overall = []
    for row in rows[header + 1:]:
        name = clean(row[0])
        if not name:
            if overall:
                break
            continue
        if name.lower() == "total":
            break
        overall.append({
            "name": name,
            "target": number(row[1]),
            "recharge": number(row[2]),
            "cards": 0.0,
            "yesterday": number(row[6] if len(row) > 6 else 0),
        })
    by_name = {row["name"].lower(): row for row in overall}
    for detail in details:
        target = by_name.get(detail["owner"].lower()) or by_name.get("others") or by_name.get("upay")
        if target:
            target["cards"] += detail["cards"]
    return overall


def main():
    source, destination = map(Path, sys.argv[1:3])
    book = load_workbook(source, read_only=True, data_only=True)
    periods = []
    for month, (summary_name, daily_name) in MONTHS.items():
        summary, daily = book[summary_name], book[daily_name]
        daily_reports = parse_daily_reports(daily)
        last_daily = daily_reports[-1]["details"] if daily_reports else []
        details = final_details(daily, last_daily)
        overall = parse_overall(summary, details)
        end = daily_reports[-1]["date"] if daily_reports else f"2026-{month:02d}-{calendar.monthrange(2026, month)[1]:02d}"
        periods.append({
            "id": f"2026-{month:02d}",
            "label": f"2026 年 {month} 月",
            "start": f"2026-{month:02d}-01",
            "end": end,
            "overall": overall,
            "details": details,
            "daily": daily_reports,
        })
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps({
        "source": "UW每日数据.xlsx",
        "metric": "consumption",
        "periods": periods,
    }, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


if __name__ == "__main__":
    main()
