"""Create the compact dashboard dataset from the monthly tabs in UB每日数据.xlsx.

Usage: python3 scripts/import_ub_excel.py /path/to/UB每日数据.xlsx app/dashboard-data.json
"""
from __future__ import annotations

import calendar
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


MONTHS = {
    1: ("1月目标进展", "代理日汇总1"), 2: ("2月目标进展", "代理日汇总2"),
    3: ("3月目标进展", "代理商日汇总3"), 4: ("4月汇总", "代理商日汇总4"),
    5: ("5月汇总", "代理日汇总5"), 6: ("6月汇总", "代理日汇总6"),
    7: ("7月汇总", "代理日汇总7"), 8: ("8月汇总", "代理日汇总8"),
    9: ("9月汇总", "代理日汇总9"),
}


def number(value):
    return float(value or 0) if isinstance(value, (int, float)) else 0.0


def label_date(value):
    return value.strftime("%Y-%m-%d") if isinstance(value, datetime) else None


def parse_daily_reports(sheet):
    """Read every dated report block from a monthly daily-summary worksheet.

    A sheet contains one repeated table per day.  If a date was exported twice,
    retain the last block rather than adding it twice.
    """
    reports = {}
    current_date, current = None, {}
    title = re.compile(r"Performance Analysis\s+(\d{4}/\d{2}/\d{2})")

    def save_current():
        if current_date:
            reports[current_date] = list(current.values())

    for row in sheet.iter_rows(values_only=True):
        first = row[0] if row else None
        match = title.search(str(first or ""))
        if match:
            save_current()
            current_date = match.group(1).replace("/", "-")
            current = {}
            continue
        if not current_date:
            continue
        owner, category, name = row[:3]
        category = str(category or "").upper()
        if not (owner and name and category in {"API", "AGENT"}):
            continue
        key = (str(owner).upper(), category, str(name).strip())
        current[key] = {
            "name": key[2], "owner": key[0],
            "type": "API" if category == "API" else "代理商",
            "recharge": number(row[5] if len(row) > 5 else 0),
            "consumption": number(row[6] if len(row) > 6 else 0),
            "cards": number(row[4] if len(row) > 4 else 0),
        }
    save_current()
    return [{"date": date, "details": reports[date]} for date in sorted(reports)]


def main():
    source, destination = map(Path, sys.argv[1:3])
    book = load_workbook(source, read_only=True, data_only=True)
    periods = []
    for month, (summary_name, daily_name) in MONTHS.items():
        summary, daily = book[summary_name], book[daily_name]
        daily_reports = parse_daily_reports(daily)
        daily_map, yesterday = {}, {}
        for row in (daily_reports[-1]["details"] if daily_reports else []):
            key = (row["owner"], row["name"])
            daily_map[key] = "API" if row["type"] == "API" else "AGENT"
            yesterday[key] = {"recharge": row["recharge"], "cards": row["cards"]}

        rows = list(summary.iter_rows(min_row=1, max_row=summary.max_row, values_only=True))
        target_row = next((i for i, r in enumerate(rows) if str(r[0] or "").strip().lower() == "maintainer" and str(r[1] or "").strip().lower() == "target"), None)
        overall = []
        if target_row is not None:
            for row in rows[target_row + 1:]:
                name = row[0]
                if not name:
                    if overall: break
                    continue
                if not isinstance(name, str): continue
                if name.strip().lower() == "total": continue
                overall.append({"name": name.strip(), "target": number(row[1]), "recharge": number(row[2]), "cards": number(row[6] if len(row) > 6 else 0), "yesterday": number(row[7] if len(row) > 7 else 0)})

        details = []
        # Apr-Sep: cumulative detail begins in the second table (around column T).
        header = rows[0] if rows else []
        recharge_col = next((i for i, value in enumerate(header) if i >= 19 and str(value or "").strip().lower() == "recharge"), None)
        if recharge_col is not None:
            maintainer_col = max(i for i, value in enumerate(header[:recharge_col]) if str(value or "").strip().lower() == "maintainer")
            agent_col = max(i for i, value in enumerate(header[:recharge_col]) if str(value or "").strip().lower() == "agent")
            for row in rows[1:]:
                owner, category, name = row[maintainer_col:maintainer_col + 3]
                if not (owner and category and name): continue
                category = str(category).upper()
                if category not in {"API", "AGENT"}: continue
                key = (str(owner).upper(), str(name).strip())
                cards = sum(number(row[i]) for i in range(agent_col + 1, recharge_col) if "open card" in str(header[i] or "").lower())
                details.append({"name": key[1], "owner": key[0], "type": "API" if category == "API" else "代理商", "recharge": number(row[recharge_col]), "consumption": number(row[recharge_col + 1] if len(row) > recharge_col + 1 else 0), "cards": cards, "yesterday": yesterday.get(key, {}).get("recharge", 0)})
        else:
            # Jan-Mar: the cumulative table starts at J and category comes from the daily tab.
            for row in rows[1:]:
                owner, name = row[9:11]
                if not (owner and name): continue
                key = (str(owner).upper(), str(name).strip())
                category = daily_map.get(key)
                if not category: continue
                details.append({"name": key[1], "owner": key[0], "type": "API" if category == "API" else "代理商", "recharge": number(row[13]), "consumption": number(row[14] if len(row) > 14 else 0), "cards": number(row[11]) + number(row[12]), "yesterday": yesterday.get(key, {}).get("recharge", 0)})

        end = f"2026-{month:02d}-{calendar.monthrange(2026, month)[1]:02d}"
        if month == 9 and len(rows) > 2 and isinstance(rows[2][0], datetime): end = label_date(rows[2][0]) or end
        if daily_reports:
            end = daily_reports[-1]["date"]
        periods.append({"id": f"2026-{month:02d}", "label": f"2026 年 {month} 月", "start": f"2026-{month:02d}-01", "end": end, "overall": overall, "details": details, "daily": daily_reports})
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps({"source": "UB每日数据.xlsx", "periods": periods}, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


if __name__ == "__main__":
    main()
