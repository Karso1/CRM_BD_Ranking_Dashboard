#!/usr/bin/env python3
"""Download five daily UPB exports and one manual-card snapshot. Never publish data."""
from __future__ import annotations

import argparse
from datetime import date, datetime, timedelta, timezone
import fcntl
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import sys
import time

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
STATE = HERE / '.local'
DATA = REPO / 'tools/daily-operations/03-原始数据/UPB每日数据'
SOURCES = {
    'cards': ('开卡', 'cardManage/cardApplyHistory', '申请时间', '开卡', ('订单号', '申请时间', '卡状态', '开卡费')),
    'recharge': ('充值', 'cardChannels/mortgageOrder', '创建时间', '充值数据', ('订单号', '创建时间', '充值金额', '状态')),
    'passto': ('Passto', 'cardChannels/passto-detailed-account', '更新时间', '消费数据', ('更新时间', '订单类型', '交易金额')),
    'reap': ('Reap', 'cardChannels/reap-channel-account-details', '更新时间', '消费数据', ('更新时间', '订单类型', '交易金额')),
    # StraitsX only offers a transaction-time search control/export column.
    'straitsx': ('StraitsX', 'cardChannels/straix-account-details', '交易时间', '消费数据', ('订单类型', '交易时间')),
}
CHANNEL_SOURCES = frozenset({'passto', 'reap', 'straitsx'})
MANUAL_ROUTE = 'cardManage/cardOpening'
MANUAL_STATUS = '激活/绑定/开卡成功'
MANUAL_SUCCESS = frozenset({'激活成功', '绑定成功', '开卡成功'})
MANUAL_REQUIRED = ('历史代理商', '用户卡片ID', '渠道卡ID', '申请日期', '订单状态')


def channel_policy(source: str) -> str:
    return f'主订单+{SOURCES[source][2]}-v1'
READY_LABELS = {
    'cards': '申请时间', 'recharge': '创建时间',
    'passto': '客户编号', 'reap': '入账状态', 'straitsx': '交易唯一ID',
}
# Some backend exports stop at 20,000 rows. Keep each export below that
# smallest known cap, including channels with a higher cap.
SAFE_EXPORT_ROWS = 18_000


class ExportButtonMissing(RuntimeError):
    """The result page rendered but did not expose its own export action."""


class PageNotReady(RuntimeError):
    """A backend search/reset did not finish, so its result cannot be trusted."""


def yesterday() -> date:
    return datetime.now(timezone.utc).date() - timedelta(days=1)


def digest(path: Path) -> str:
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def parse_export_time(value: object, source: str, row_number: int) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    try:
        return datetime.fromisoformat(str(value).strip().replace('/', '-')).replace(tzinfo=None)
    except ValueError as exc:
        raise ValueError(f'{source} 第 {row_number} 行日期无法识别') from exc


def validate(path: Path, source: str, day: date, start: datetime | None = None, end: datetime | None = None) -> int:
    from openpyxl import load_workbook
    required = SOURCES[source][4]
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book:
            # Some backend exports declare A1 as the used range even when the
            # sheet contains thousands of rows and dozens of columns.
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            header = tuple(str(v).strip() if v is not None else '' for v in next(rows, ()))
            if not set(required).issubset(header):
                continue
            time_index = header.index(SOURCES[source][2]) if SOURCES[source][2] in header else None
            if time_index is None:
                raise ValueError(f'{source} 文件缺少{SOURCES[source][2]}，无法核验下载日期。')
            order_index = header.index('订单类型') if source in CHANNEL_SOURCES else None
            count = 0
            for row in rows:
                if not any(v is not None for v in row):
                    continue
                if order_index is not None and str(row[order_index]).strip() != '主订单':
                    raise ValueError(f'{source} 第 {count + 2} 行不是主订单，未导入。')
                actual = parse_export_time(row[time_index], source, count + 2)
                if actual.date() != day:
                    raise ValueError(f'{source} 包含{SOURCES[source][2]} {actual.date()} 的数据，预期仅 {day}，未导入。')
                if (start is not None and actual < start) or (end is not None and actual > end):
                    raise ValueError(f'{source} 包含区间外{SOURCES[source][2]} {actual}，预期 {start} 至 {end}，未导入。')
                count += 1
            return count
        raise ValueError(f'{source} 文件缺少必需表头：{required}')
    finally:
        book.close()


def validate_manual(path: Path) -> int:
    """The full snapshot must be usable for attribution and contain successes only."""
    from openpyxl import load_workbook
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book:
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            header = tuple(str(value).strip() if value is not None else '' for value in next(rows, ()))
            if not set(MANUAL_REQUIRED).issubset(header):
                continue
            status_index = header.index('订单状态')
            card_index = header.index('用户卡片ID')
            count = 0
            for row in rows:
                if not any(value is not None for value in row):
                    continue
                count += 1
                if str(row[status_index] or '').strip() not in MANUAL_SUCCESS:
                    raise ValueError(f'手动开卡第 {count + 1} 行不是成功状态，未替换旧表。')
                if not str(row[card_index] or '').strip():
                    raise ValueError(f'手动开卡第 {count + 1} 行缺少用户卡片ID，未替换旧表。')
            if count == 0:
                raise ValueError('手动开卡导出为空，未替换旧表。')
            return count
        raise ValueError(f'手动开卡文件缺少必需表头：{MANUAL_REQUIRED}')
    finally:
        book.close()


def set_datetime(page, placeholder: str, index: int, value: str, label: str | None = None) -> None:
    if label is None:
        field = page.get_by_placeholder(placeholder, exact=True).nth(index)
    else:
        item = page.locator('main .ant-form-item').filter(has=page.locator(f'label[title="{label}"]'))
        field = item.get_by_placeholder(placeholder, exact=True)
        if field.count() != 1:
            raise RuntimeError(f'{label} 的{placeholder}控件未找到，已停止导出。')
    for attempt in range(3):
        page.keyboard.press('Escape')
        field.click()  # The site's date inputs become editable only after focus.
        for _ in range(3):
            if field.is_editable():
                break
            page.wait_for_timeout(150)
            field.click()
        if not field.is_editable():
            raise RuntimeError(f'日期控件未进入编辑状态：{placeholder}')
        field.fill(value)
        confirm = page.get_by_role('button', name=re.compile(r'^确\s*定$'))
        confirmed = False
        for button in confirm.all():
            if button.is_visible():
                button.click()
                confirmed = True
                break
        if not confirmed:
            field.press('Enter')
        page.wait_for_timeout(350)
        if field.input_value() == value:
            page.keyboard.press('Escape')
            return
        page.wait_for_timeout(400)
    raise RuntimeError(f'日期条件未生效：期望 {value}，实际 {field.input_value()!r}')


def wait_idle(page) -> None:
    page.locator('main .ant-spin-blur').wait_for(state='hidden', timeout=60000)


def wait_controls_ready(page, phase: str) -> None:
    """Do not search/export while an Ant Design action is still spinning."""
    try:
        page.locator('main .ant-btn-loading').wait_for(state='hidden', timeout=60000)
        wait_idle(page)
    except Exception as exc:
        raise PageNotReady(f'{phase} 超过 60 秒仍在加载；将刷新页面并重新核对筛选。') from exc


def scroll_to_top(page) -> None:
    """SPA routes retain the prior page scroll, hiding their export toolbar."""
    page.evaluate('''() => {
        window.scrollTo(0, 0);
        const main = document.querySelector('main');
        if (main) main.scrollTop = 0;
        for (const area of document.querySelectorAll('main *')) {
            if (area.scrollTop > 0 && area.scrollHeight > area.clientHeight) area.scrollTop = 0;
        }
    }''')


def manual_page_count(page) -> int:
    """Read the table pagination, not an unrelated count elsewhere in main."""
    totals = page.locator('main .pagination-container > p')
    if totals.count() != 1:
        raise RuntimeError('手动开卡分页计数未找到或不唯一，停止导出。')
    match = re.fullmatch(r'共\s*([\d,]+)\s*条', totals.inner_text().strip())
    if not match:
        raise RuntimeError('手动开卡分页计数格式已变化，停止导出。')
    return int(match.group(1).replace(',', ''))


def wait_manual_controls_ready(page) -> None:
    """Reset/search buttons and the status menu must finish before querying."""
    reset = page.get_by_role('button', name=re.compile('重置'))
    reset.locator('span.ant-btn-loading-icon').wait_for(state='hidden', timeout=60000)
    page.locator('main .ant-btn-loading').wait_for(state='hidden', timeout=60000)
    wait_idle(page)


def manual_visible_rows_success(page) -> bool:
    """The first successful row alone does not prove the filter was applied."""
    headers = [heading.inner_text().strip() for heading in page.locator('main .ant-table-thead th').all()]
    if '卡片状态' not in headers:
        raise RuntimeError('手动开卡列表缺少卡片状态列，停止导出。')
    index = headers.index('卡片状态')
    rows = page.locator('main .ant-table-tbody tr:visible:not([aria-hidden="true"]):not(.ant-table-placeholder)').all()
    return bool(rows) and all(row.locator('td').nth(index).inner_text().strip() in MANUAL_SUCCESS for row in rows)


def manual_result_count_after_search(page) -> int:
    """Never treat the transient zero shown during a pending search as final."""
    search = page.get_by_role('button', name=re.compile('搜索'))
    for attempt in range(3):
        page.keyboard.press('Escape')
        page.locator('.ant-select-dropdown:visible').wait_for(state='hidden', timeout=10000)
        wait_manual_controls_ready(page)
        search.click()
        wait_manual_controls_ready(page)
        page.wait_for_load_state('networkidle', timeout=60000)
        wait_manual_controls_ready(page)
        # A zero row count may be a stale/transient response. Give the table
        # time to render, then explicitly query again instead of exporting it.
        page.wait_for_timeout(1200)
        count = manual_page_count(page)
        if count > 0 and manual_visible_rows_success(page):
            return count
        if attempt < 2:
            print(f'手动开卡查询未呈现完整的成功状态筛选，等待后重新搜索（{attempt + 1}/2）…', flush=True)
            page.wait_for_timeout(1500)
    raise PageNotReady('手动开卡连续 3 次搜索仍未显示成功状态的列表；旧表保持不变。')


def prepare_page(page, source: str, start: datetime, end: datetime) -> int | None:
    _, route, _, _, _ = SOURCES[source]
    page.goto('https://admin.zfuucr.com/#/' + route)
    search = page.get_by_role('button', name=re.compile('搜索'))
    for attempt in range(5):
        try:
            search.wait_for(timeout=8000)
            break
        except Exception:
            if page.get_by_placeholder('请输入登录账号').count():
                if attempt > 0:
                    raise RuntimeError('后台登录未完成。')
                print('请在自动打开的 Chrome 窗口中登录业务后台；完成后回到此窗口按 Enter。', flush=True)
                input()
            elif '#/home' not in page.url and '#/dashboard' not in page.url and '#/index' not in page.url and '#/' + route in page.url:
                # Route is correct but still loading: give the page another chance.
                continue
            # The site redirects to its home route while restoring login.
            page.goto('https://admin.zfuucr.com/#/' + route)
    else:
        raise RuntimeError(f'{source} 页面未能打开。')
    page.locator('main').get_by_text(READY_LABELS[source], exact=False).first.wait_for(timeout=30000)
    wait_controls_ready(page, f'{source} 页面')
    scroll_to_top(page)
    # The async-export history drawer can stay open when the same route is
    # revisited for the next time slice and intercept the date controls.
    if history_open(page):
        close_history(page)
    # A different timezone would change which business day is exported.
    if not page.get_by_text('UTC+0', exact=True).is_visible():
        raise RuntimeError('后台时区不是 UTC+0，请先在后台切换为 UTC+0 后重试。')
    page.get_by_role('button', name=re.compile('重置')).click()
    wait_controls_ready(page, f'{source} 重置')
    page.wait_for_timeout(1000)
    # Clear every default date filter, especially platform posting time.
    clear = page.get_by_role('button', name='close-circle', exact=True)
    for _ in range(12):
        if not clear.count():
            break
        clear.first.click()
    if clear.count():
        raise RuntimeError('默认日期筛选未清除，已停止导出。')
    if source in CHANNEL_SOURCES:
        order = page.locator('#form_item_orderType')
        if order.count() != 1:
            raise RuntimeError(f'{source} 订单类型控件未找到，停止导出。')
        selected = order.locator('xpath=../..')  # input -> search span -> selector
        if '主订单' not in selected.inner_text():
            selected.click()
            page.locator('.ant-select-dropdown:visible').get_by_text('主订单', exact=True).click()
        if '主订单' not in selected.inner_text():
            raise RuntimeError(f'{source} 订单类型未设为主订单，停止导出。')
    date_label = SOURCES[source][2] if source in CHANNEL_SOURCES else None
    set_datetime(page, '开始时间', 0, start.strftime('%Y-%m-%d %H:%M:%S'), date_label)
    set_datetime(page, '结束时间', 0, end.strftime('%Y-%m-%d %H:%M:%S'), date_label)
    search.click()
    page.wait_for_timeout(1500)
    wait_controls_ready(page, f'{source} 搜索')
    scroll_to_top(page)
    text = page.locator('main').inner_text()
    total = re.search(r'共\s*([\d,]+)\s*条', text)
    return int(total.group(1).replace(',', '')) if total else None


def history_tasks(page) -> dict[str, object]:
    cards = page.locator('div.mb-3.rounded-lg.border').filter(visible=True)
    result = {}
    for card in cards.all():
        match = re.search(r'任务ID[：:]\s*(\d+)', card.inner_text())
        if match:
            result[match.group(1)] = card
    return result


def select_new_task(before: set[str], current: set[str]) -> str | None:
    """Task IDs are time-ordered; older delayed tasks may join the list too."""
    new = current - before
    return max(new, key=int) if new else None


def history_open(page) -> bool:
    return page.get_by_text(re.compile('最近 20 个导出历史')).filter(visible=True).count() > 0


def close_history(page) -> None:
    page.get_by_role('button', name='Close', exact=True).filter(visible=True).first.click()
    page.get_by_text(re.compile('最近 20 个导出历史')).filter(visible=True).first.wait_for(state='hidden')


def ensure_history(page, source: str) -> None:
    route = SOURCES[source][1]
    if page.get_by_placeholder('请输入登录账号').count():
        print('后台会话已过期。请在 Chrome 窗口重新登录，然后回到此窗口按 Enter。', flush=True)
        input()
        page.goto('https://admin.zfuucr.com/#/' + route)
    history = page.get_by_role('button', name=re.compile('导出历史'))
    if not history.count():
        page.goto('https://admin.zfuucr.com/#/' + route)
        history.wait_for(timeout=30000)
    if not history_open(page):
        history.click()
    page.get_by_text(re.compile('最近 20 个导出历史')).filter(visible=True).first.wait_for()


def find_export_button(page, source: str):
    """Exclude the history button; some pages render Export only after search."""
    scroll_to_top(page)
    deadline = time.monotonic() + 12
    while time.monotonic() < deadline:
        candidates = page.get_by_role('button', name=re.compile(r'导\s*出')).filter(visible=True)
        for candidate in candidates.all():
            label = candidate.inner_text().strip()
            if '历史' not in label and '记录' not in label and re.search(r'导\s*出', label):
                return candidate
        page.wait_for_timeout(300)
    labels = [button.inner_text().strip() for button in page.get_by_role('button').filter(visible=True).all()]
    raise ExportButtonMissing(f'{source} 筛选结果已显示，但页面未出现导出按钮；当前可见按钮：{labels[:18]}')


def export_file(page, source: str, destination: Path) -> tuple[str, int | None]:
    history = page.get_by_role('button', name=re.compile('导出历史'))
    if not history.count():
        task_rows = None
        export = find_export_button(page, source)
        with page.expect_download(timeout=180000) as event:
            export.click()
        download = event.value
    else:
        ensure_history(page, source)
        refresh = page.get_by_role('button', name=re.compile('刷新'))
        if refresh.count() and refresh.is_visible():
            refresh.click()
        page.wait_for_timeout(1500)
        before = set(history_tasks(page))
        close_history(page)
        export = find_export_button(page, source)
        export.click()
        # Export tasks are asynchronous; never select an arbitrary old file.
        page.wait_for_timeout(2000)
        ensure_history(page, source)
        page.wait_for_timeout(1200)
        deadline = time.monotonic() + 600
        task_id = None
        while time.monotonic() < deadline:
            ensure_history(page, source)
            tasks = history_tasks(page)
            if task_id is None:
                task_id = select_new_task(before, set(tasks))
            if task_id and task_id in tasks:
                card = tasks[task_id]
                text = card.inner_text()
                if '已失败' in text:
                    raise RuntimeError(f'{source} 后台导出失败，任务 {task_id}')
                button = card.get_by_role('button', name=re.compile('下载'))
                if '已成功' in text and button.count():
                    row_match = re.search(r'([\d,]+)\s*行', text)
                    if not row_match:
                        raise RuntimeError(f'{source} 导出任务 {task_id} 未显示行数，无法核验完整性。')
                    task_rows = int(row_match.group(1).replace(',', ''))
                    with page.expect_download(timeout=180000) as event:
                        button.click()
                    download = event.value
                    break
            page.wait_for_timeout(3000)
            refresh = page.get_by_role('button', name=re.compile('刷新'))
            if refresh.count() and refresh.is_visible():
                refresh.click()
            else:
                # Some channel pages replace the history drawer while a task
                # is pending; reopening it asks the backend for fresh tasks.
                if history_open(page):
                    close_history(page)
                ensure_history(page, source)
        else:
            raise RuntimeError(f'{source} 导出超过 10 分钟，未导入。')
    download.save_as(destination)
    if download.failure():
        raise RuntimeError(f'{source} 下载失败')
    return download.suggested_filename, task_rows


def split_window(start: datetime, end: datetime) -> tuple[tuple[datetime, datetime], tuple[datetime, datetime]]:
    """Partition an inclusive UTC-second interval without overlap or gaps."""
    seconds = int((end - start).total_seconds())
    if seconds < 1:
        raise RuntimeError(f'{start} 单秒数据仍超过安全导出行数，无法再拆分；未导入。')
    middle = start + timedelta(seconds=seconds // 2)
    return (start, middle), (middle + timedelta(seconds=1), end)


def download_window(page, source: str, day: date, start: datetime, end: datetime, batch: Path) -> tuple[list[tuple[Path, int]], int]:
    """Split only verified large exports; the page's total can lag a search."""
    expected = None
    def subdivide() -> tuple[list[tuple[Path, int]], int]:
        left, right = split_window(start, end)
        print(f'{SOURCES[source][0]} {start:%Y-%m-%d %H:%M:%S} 至 {end:%H:%M:%S} 导出接近后台上限，拆成两段。', flush=True)
        first_parts, first_count = download_window(page, source, day, *left, batch)
        second_parts, second_count = download_window(page, source, day, *right, batch)
        return first_parts + second_parts, first_count + second_count

    destination = batch / f'{source}_{start:%Y%m%d%H%M%S}_{end:%H%M%S}.xlsx'
    for attempt in range(2):
        try:
            expected = prepare_page(page, source, start, end)
            _, task_rows = export_file(page, source, destination)
            break
        except (ExportButtonMissing, PageNotReady) as exc:
            if attempt:
                raise
            print(f'{SOURCES[source][0]} 页面未就绪（{exc}），刷新后重新核对 {start:%Y-%m-%d %H:%M:%S} 至 {end:%H:%M:%S} 的筛选并重试一次…', flush=True)
            page.reload(wait_until='domcontentloaded')
    count = validate(destination, source, day, start, end)
    if task_rows is not None and count != task_rows:
        raise RuntimeError(f'{source} 导出任务显示 {task_rows} 行，文件实际 {count} 行，未导入。')
    if count >= SAFE_EXPORT_ROWS:
        destination.unlink()
        return subdivide()
    if task_rows is None and expected is None:
        raise RuntimeError(f'{source} 页面未显示总条数，且没有导出任务行数，无法排除截断。')
    if task_rows is None and expected >= SAFE_EXPORT_ROWS:
        destination.unlink()
        return subdivide()
    if task_rows is None and source != 'straitsx' and count != expected:
        raise RuntimeError(f'{source} 页面 {expected} 条，导出 {count} 行，未导入。')
    if source == 'straitsx' and count != expected:
        raise RuntimeError(
            f'StraitsX 页面 {expected} 条，导出 {count} 行，数量不符，未导入。'
        )
    if count == 0 and task_rows is None and expected != 0:
        raise RuntimeError(f'{source} 导出为空，但页面未明确显示 0 条。')
    return [(destination, count)], count


def merge_parts(source: str, day: date, parts: list[tuple[Path, int]], destination: Path) -> int:
    from openpyxl import Workbook, load_workbook
    if not parts:
        raise RuntimeError(f'{source} 没有可合并的导出文件。')
    if len(parts) == 1:
        shutil.copy2(parts[0][0], destination)
    else:
        output = Workbook(write_only=True)
        sheet_out = output.create_sheet('sheet1')
        expected_header = None
        for path, _ in parts:
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                sheet = book.active
                sheet.reset_dimensions()
                rows = sheet.iter_rows(values_only=True)
                header = tuple(next(rows, ()))
                if expected_header is None:
                    expected_header = header
                    sheet_out.append(header)
                elif header != expected_header:
                    raise RuntimeError(f'{source} 分段导出表头不一致，未合并。')
                for row in rows:
                    if any(value is not None for value in row):
                        sheet_out.append(row)
            finally:
                book.close()
        output.save(destination)
    count = validate(destination, source, day)
    if count != sum(part_count for _, part_count in parts):
        raise RuntimeError(f'{source} 合并行数不等于各分段行数之和，未导入。')
    return count


def prepare_manual_page(page) -> int:
    page.goto('https://admin.zfuucr.com/#/' + MANUAL_ROUTE)
    status = page.locator('#form_item_cardState')
    for attempt in range(5):
        try:
            status.wait_for(timeout=8000)
            break
        except Exception:
            if page.get_by_placeholder('请输入登录账号').count():
                if attempt > 0:
                    raise RuntimeError('后台登录未完成。')
                print('请在自动打开的 Chrome 窗口中登录业务后台；完成后回到此窗口按 Enter。', flush=True)
                input()
            page.goto('https://admin.zfuucr.com/#/' + MANUAL_ROUTE)
    else:
        raise RuntimeError('手动开卡页面未能打开。')
    wait_idle(page)
    page.get_by_role('button', name=re.compile('重置')).click()
    wait_manual_controls_ready(page)
    selected = status.locator('xpath=../..')
    if MANUAL_STATUS not in selected.inner_text():
        selected.click()
        page.locator('.ant-select-dropdown:visible').get_by_text(MANUAL_STATUS, exact=True).click()
    if MANUAL_STATUS not in selected.inner_text():
        raise RuntimeError('手动开卡状态筛选未设为“激活/绑定/开卡成功”，停止导出。')
    count = manual_result_count_after_search(page)
    if MANUAL_STATUS not in selected.inner_text():
        raise RuntimeError('手动开卡搜索后成功状态筛选丢失；旧表保持不变。')
    if count >= SAFE_EXPORT_ROWS:
        raise RuntimeError(f'手动开卡页面显示 {count} 条，不能确认后台会完整导出；旧表保持不变。')
    return count


def download_manual(page, batch: Path) -> tuple[Path, int, str]:
    destination = batch / 'manual_cards.xlsx'
    for attempt in range(2):
        try:
            expected = prepare_manual_page(page)
            with page.expect_download(timeout=180000) as event:
                page.get_by_role('button', name=re.compile(r'^导\s*出$')).click()
            download = event.value
            download.save_as(destination)
            if download.failure():
                raise RuntimeError('手动开卡下载失败，旧表保持不变。')
            actual = validate_manual(destination)
            # A delayed search can leave the pre-export pagination stale.
            latest = manual_page_count(page)
            if latest != expected:
                expected = latest
            if actual < expected or actual > expected + 10:
                raise PageNotReady(f'手动开卡页面 {expected} 条，导出 {actual} 行，数量不一致；旧表保持不变。')
            if actual > expected:
                print(f'手动开卡列表显示 {expected} 条，导出为 {actual} 行；已逐行核验成功状态及卡 ID。', flush=True)
            return destination, actual, download.suggested_filename
        except PageNotReady as exc:
            if attempt:
                raise
            print(f'{exc} 刷新手动开卡页面后重新筛选并导出一次…', flush=True)
            page.reload(wait_until='domcontentloaded')
    raise AssertionError('unreachable')


def replace_manual_snapshot(export: Path, data: Path, state: Path, rows: int, filename: str) -> tuple[Path, Path | None]:
    """Replace the sole original snapshot after backing it up and verifying it."""
    if validate_manual(export) != rows:
        raise RuntimeError('手动开卡暂存文件校验失败，旧表保持不变。')
    folder = data / '手动开卡'
    if not folder.is_dir():
        raise RuntimeError(f'手动开卡原始数据目录不存在：{folder}')
    originals = [path for path in folder.glob('*.xlsx') if not path.name.startswith('~$')]
    if len(originals) > 1:
        raise RuntimeError('手动开卡原始数据目录存在多份 Excel，无法确认替换哪一份。')
    previous = originals[0] if originals else None
    target = folder / '手动开卡-最新.xlsx'
    backup = None
    if previous:
        backup_dir = state / 'manual-backups' / datetime.now().strftime('%Y%m%d-%H%M%S-%f')
        backup_dir.mkdir(parents=True)
        backup = backup_dir / previous.name
        shutil.copy2(previous, backup)
        if digest(backup) != digest(previous):
            raise RuntimeError('手动开卡旧文件备份校验失败，未替换。')
    staged = target.with_name(target.stem + '.manual-pending.xlsx')
    if staged.exists():
        raise RuntimeError(f'{staged} 已存在，请先核查上次更新记录。')
    installed = False
    removed_legacy = False
    try:
        shutil.copy2(export, staged)
        if digest(staged) != digest(export) or validate_manual(staged) != rows:
            raise RuntimeError('手动开卡新文件暂存校验失败，未替换。')
        os.replace(staged, target)
        installed = True
        if previous is not None and previous != target:
            previous.unlink()
            removed_legacy = True
        record = {
            'updated_at': datetime.now(timezone.utc).isoformat(),
            'path': str(target), 'sha256': digest(target),
            'rows': rows, 'export_filename': filename,
            'backup': str(backup) if backup else None,
        }
        pending = state / 'manual-snapshot.pending'
        pending.write_text(json.dumps(record, ensure_ascii=False, indent=2))
        os.replace(pending, state / 'manual-snapshot.json')
    except BaseException:
        staged.unlink(missing_ok=True)
        (state / 'manual-snapshot.pending').unlink(missing_ok=True)
        if installed:
            if previous == target and backup:
                rollback = target.with_name(target.stem + '.manual-rollback-pending.xlsx')
                shutil.copy2(backup, rollback)
                os.replace(rollback, target)
            else:
                target.unlink(missing_ok=True)
                if removed_legacy and backup:
                    shutil.copy2(backup, previous)
        raise
    return target, backup


def commit_batch(batch: Path, day: date, records: dict, data: Path) -> None:
    """Prepare every file before exposing the batch; rollback on any error."""
    planned = []
    for source, record in records.items():
        folder = data / SOURCES[source][3] / '每日自动下载' / day.strftime('%Y-%m')
        target = folder / f'{SOURCES[source][0]}_{day}.xlsx'
        if target.exists():
            raise RuntimeError(f'目标文件已存在，未覆盖：{target}')
        planned.append((batch / f'{source}.xlsx', target, record))
    installed = []
    staged_files = []
    try:
        for source_path, target, record in planned:
            target.parent.mkdir(parents=True, exist_ok=True)
            staged = target.with_suffix('.pending')
            staged_files.append(staged)
            shutil.copy2(source_path, staged)
            os.replace(staged, target)
            installed.append(target)
            record['path'] = str(target)
            record['sha256'] = digest(target)
    except BaseException:
        for target in installed:
            target.unlink()  # Only newly created copies; originals stay in batch.
        for staged in staged_files:
            staged.unlink(missing_ok=True)
        raise


def repair_batch(batch: Path, day: date, records: dict, data: Path, state: Path) -> Path:
    """Replace only the three channel files, keeping restorable originals."""
    manifest = state / f'{day}.json'
    old = json.loads(manifest.read_text())
    if not completed_batch(day, state):
        raise RuntimeError(f'{day} 没有可修复的完整旧批次。')
    backup = state / 'repairs' / f'{day}-{datetime.now():%H%M%S%f}'
    backup.mkdir(parents=True)
    shutil.copy2(manifest, backup / manifest.name)
    planned = []
    for source in CHANNEL_SOURCES:
        target = data / SOURCES[source][3] / '每日自动下载' / day.strftime('%Y-%m') / f'{SOURCES[source][0]}_{day}.xlsx'
        if Path(old['files'][source]['path']) != target:
            raise RuntimeError(f'{source} 旧文件不在预期数据目录，停止修复。')
        replacement = batch / f'{source}.xlsx'
        if not replacement.is_file():
            raise RuntimeError(f'{source} 修复文件缺失，停止替换。')
        if validate(replacement, source, day) != records[source]['rows']:
            raise RuntimeError(f'{source} 修复文件行数核验失败。')
        saved = backup / target.name
        shutil.copy2(target, saved)
        if digest(saved) != old['files'][source]['sha256']:
            raise RuntimeError(f'{source} 旧文件备份校验失败，停止替换。')
        staged = target.with_suffix('.repair-pending')
        if staged.exists():
            raise RuntimeError(f'{staged} 已存在，请先核查上次修复记录。')
        planned.append((source, replacement, target, staged, saved))
    installed = []
    try:
        for source, replacement, target, staged, saved in planned:
            shutil.copy2(replacement, staged)
            if digest(staged) != digest(replacement):
                raise RuntimeError(f'{source} 临时文件校验失败。')
        for source, replacement, target, staged, saved in planned:
            os.replace(staged, target)
            installed.append((target, saved))
            records[source]['path'] = str(target)
            records[source]['sha256'] = digest(target)
            records[source]['filter_policy'] = channel_policy(source)
        old['files'] = records
        old['repair_backup'] = str(backup)
        pending_manifest = manifest.with_suffix('.repair-pending')
        pending_manifest.write_text(json.dumps(old, ensure_ascii=False, indent=2))
        os.replace(pending_manifest, manifest)
    except BaseException:
        for target, saved in installed:
            rollback = target.with_suffix('.rollback-pending')
            shutil.copy2(saved, rollback)
            os.replace(rollback, target)
        for _, _, _, staged, _ in planned:
            staged.unlink(missing_ok=True)
        manifest.with_suffix('.repair-pending').unlink(missing_ok=True)
        raise
    return backup


def cached_download(source: str, day: date, batch: Path) -> tuple[Path, int] | None:
    """Resume an interrupted same-day download from already validated files."""
    folders = sorted((STATE / 'batches').glob(f'{day}-*'), reverse=True)
    for folder in folders:
        if folder == batch:
            continue
        candidate = folder / f'{source}.xlsx'
        if not candidate.is_file():
            continue
        if source in CHANNEL_SOURCES:
            marker = folder / f'{source}.policy'
            if not marker.is_file() or marker.read_text().strip() != channel_policy(source):
                continue
        try:
            count = validate(candidate, source, day)
        except (OSError, ValueError):
            continue
        if count < SAFE_EXPORT_ROWS:
            return candidate, count
    return None


def dates_to_download(single: date | None, start: date | None, end: date | None, today: date, state: Path) -> list[date]:
    """Default to every day since the first completed daily batch, including gaps."""
    if single is not None and (start is not None or end is not None):
        raise ValueError('指定单日和日期范围不能同时使用。')
    if (start is None) != (end is None):
        raise ValueError('日期范围需要同时填写开始和结束日期。')
    yesterday_utc = today - timedelta(days=1)
    if single is not None:
        first = last = single
    elif start is not None and end is not None:
        first, last = start, end
    else:
        completed = []
        for manifest in state.glob('????-??-??.json'):
            try:
                completed.append(date.fromisoformat(manifest.stem))
            except ValueError:
                continue
        first, last = min(completed, default=yesterday_utc), yesterday_utc
    if first > last:
        raise ValueError('开始日期不能晚于结束日期。')
    if last >= today:
        raise ValueError('只支持已经结束的 UTC 日期。')
    return [first + timedelta(days=offset) for offset in range((last - first).days + 1)]


def completed_batch(day: date, state: Path) -> bool:
    manifest = state / f'{day}.json'
    if not manifest.exists():
        return False
    try:
        saved = json.loads(manifest.read_text())
        records = saved['files']
        if saved['date'] != str(day) or set(records) != set(SOURCES):
            raise ValueError('记录中的日期或文件类别不完整')
        if all(Path(record['path']).is_file() and digest(Path(record['path'])) == record['sha256'] for record in records.values()):
            return True
    except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise RuntimeError(f'{day} 完成记录无法校验：{exc}') from exc
    raise RuntimeError(f'{day} 已完成记录对应的文件发生变化，请核查后再运行。')


def download_day(day: date, data: Path, state: Path, page, repair: bool = False) -> None:
    """Commit one day completely, then reuse the browser on the next day."""
    batch = state / 'batches' / f'{day}-{datetime.now():%H%M%S%f}'
    batch.mkdir(parents=True)
    records = {}
    pending = []
    old_records = json.loads((state / f'{day}.json').read_text())['files'] if repair else {}
    for source in SOURCES:
        if repair and source not in CHANNEL_SOURCES:
            records[source] = old_records[source].copy()
            continue
        cached = cached_download(source, day, batch)
        if cached:
            source_path, count = cached
            shutil.copy2(source_path, batch / f'{source}.xlsx')
            records[source] = {'rows': count, 'filename': f'已校验的本地暂存：{source_path.name}'}
            if source in CHANNEL_SOURCES:
                records[source]['filter_policy'] = channel_policy(source)
            print(f'{SOURCES[source][0]}：复用已校验的 {count} 行。', flush=True)
        else:
            pending.append(source)
    if pending:
        if page is None:
            raise RuntimeError('下载页面未初始化。')
        try:
            for source in pending:
                print(f'正在下载 {SOURCES[source][0]}（{day}）…', flush=True)
                target = batch / f'{source}.xlsx'
                start = datetime.combine(day, datetime.min.time())
                end = start + timedelta(days=1) - timedelta(seconds=1)
                parts, _ = download_window(page, source, day, start, end, batch)
                count = merge_parts(source, day, parts, target)
                if source in CHANNEL_SOURCES:
                    (batch / f'{source}.policy').write_text(channel_policy(source))
                records[source] = {'rows': count, 'filename': f'分段导出 {len(parts)} 个文件'}
                if source in CHANNEL_SOURCES:
                    records[source]['filter_policy'] = channel_policy(source)
                print(f'已校验：{count} 行。', flush=True)
        except BaseException:
            page.screenshot(path=str(batch / 'failure.png'), full_page=True)
            raise
    # Share the existing staging lock so a staging update cannot read a
    # partially installed five-file batch.
    with (REPO / 'tools/daily-operations/.staging.lock').open('a') as update_lock:
        fcntl.flock(update_lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        if repair:
            backup = repair_batch(batch, day, records, data, state)
        else:
            commit_batch(batch, day, records, data)
            (state / f'{day}.json').write_text(json.dumps({'date': str(day), 'files': records}, ensure_ascii=False, indent=2))
    if repair:
        print(f'{day} 三个渠道原始文件已修复；旧文件备份：{backup}', flush=True)
    else:
        print(f'{day} 五类文件全部保存成功。', flush=True)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--date', type=date.fromisoformat, help='只下载指定 UTC 日期')
    parser.add_argument('--start-date', type=date.fromisoformat, help='补下载的开始日期，须与 --end-date 一起填写')
    parser.add_argument('--end-date', type=date.fromisoformat, help='补下载的结束日期，须与 --start-date 一起填写')
    parser.add_argument('--repair-date', type=date.fromisoformat, help='重新下载已完成日期的三个渠道，先备份再替换')
    parser.add_argument('--daily-only', action='store_true', help='只补齐五类每日数据；手动开卡页面异常时用于安全续跑')
    parser.add_argument('--dry-run', action='store_true')
    parser.add_argument('--yes', action='store_true', help='确认执行超过 7 天的批量下载')
    parser.add_argument('--data-dir', type=Path, default=DATA)
    args = parser.parse_args()
    if args.repair_date is not None and (args.date is not None or args.start_date is not None or args.end_date is not None):
        parser.error('--repair-date 不能和其他日期选项同时使用。')
    if args.repair_date is not None and args.daily_only:
        parser.error('--repair-date 已不更新手动开卡，无需 --daily-only。')
    try:
        days = dates_to_download(args.repair_date or args.date, args.start_date, args.end_date, datetime.now(timezone.utc).date(), STATE)
    except ValueError as exc:
        parser.error(str(exc))
    data = args.data_dir.resolve(strict=True)
    if args.dry_run:
        print(f'日期范围：{days[0]} 至 {days[-1]}（UTC+0），共 {len(days)} 天\n数据目录：{data}')
        for day in days:
            print(f'{day}：' + ('待修复三个渠道' if args.repair_date else '已完成' if completed_batch(day, STATE) else '待下载'))
        if not args.repair_date and not args.daily_only:
            print('手动开卡：每次运行更新一次成功状态的全量快照，先备份旧表。')
        return 0
    from playwright.sync_api import sync_playwright
    os.umask(0o077)
    STATE.mkdir(parents=True, exist_ok=True)
    with (STATE / 'download.lock').open('w') as lock:
        try:
            fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            raise RuntimeError('已有下载任务正在运行，请等待它完成。')
        pending_days = days if args.repair_date else [day for day in days if not completed_batch(day, STATE)]
        if args.repair_date and not completed_batch(args.repair_date, STATE):
            raise RuntimeError(f'{args.repair_date} 没有完整旧批次，不能修复。')
        if not pending_days and not args.daily_only:
            print(f'{days[0]} 至 {days[-1]} 的五类每日数据均已保存；继续更新手动开卡全量快照。')
        preview = '、'.join(map(str, pending_days[:10]))
        if len(pending_days) > 10:
            preview += f' … {pending_days[-1]}'
        label = '待修复' if args.repair_date else '待下载'
        count = len(CHANNEL_SOURCES) if args.repair_date else len(SOURCES)
        if pending_days:
            print(f'{label} {len(pending_days)} 天（至少 {len(pending_days) * count} 个后台导出任务）：{preview}', flush=True)
        if len(pending_days) > 7 and not args.yes:
            if not sys.stdin.isatty():
                raise RuntimeError('超过 7 天的下载需要在终端输入“继续”确认，或在命令行加 --yes。')
            if input('长时间下载可能产生很多导出任务。确认开始请输入“继续”：').strip() != '继续':
                print('已取消，原始数据未改动。')
                return 0
        with sync_playwright() as p:
            # Independent profile: no copying of the user's Chrome credentials.
            context = p.chromium.launch_persistent_context(
                str(STATE / 'browser'), channel='chrome', headless=False,
                accept_downloads=True, timezone_id='UTC',
            )
            page = context.new_page()
            page.set_default_timeout(20000)
            try:
                manual = None
                if not args.repair_date and not args.daily_only:
                    manual_batch = STATE / 'manual-batches' / datetime.now().strftime('%Y%m%d-%H%M%S-%f')
                    manual_batch.mkdir(parents=True)
                    print('正在下载手动开卡成功状态的全量快照…', flush=True)
                    try:
                        manual = download_manual(page, manual_batch)
                    except BaseException:
                        page.screenshot(path=str(manual_batch / 'failure.png'))
                        raise
                    print(f'手动开卡已校验：{manual[1]} 行。', flush=True)
                for day in pending_days:
                    download_day(day, data, STATE, page, repair=bool(args.repair_date))
                if manual is not None:
                    with (REPO / 'tools/daily-operations/.staging.lock').open('a') as update_lock:
                        fcntl.flock(update_lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
                        target, backup = replace_manual_snapshot(manual[0], data, STATE, manual[1], manual[2])
                    print(f'手动开卡快照已替换：{target}（{manual[1]} 行）；旧文件备份：{backup or "无旧文件"}', flush=True)
            finally:
                context.close()
        print('所选数据全部下载完成。现在可运行 UPB-测试更新，再确认正式发布。')
    return 0


if __name__ == '__main__':
    try:
        raise SystemExit(main())
    except (Exception, KeyboardInterrupt) as error:
        print(f'未完成：{error}\n已下载文件保留在本机暂存目录，未自动发布网站。', file=sys.stderr)
        raise SystemExit(1)
